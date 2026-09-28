"""互动吧风格 Excel（.xlsx）人员导入：解析表头、下载身份识别照片并录脸。"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.config import MATCH_THRESHOLD, MAX_SAMPLES_PER_PERSON, UPLOAD_DIR
from app.face_engine import (
    FaceEngineError,
    bbox_to_json,
    encode_image_jpeg,
    embedding_to_bytes,
    face_engine,
)
from app.gallery_index import gallery_index
from app.models import FaceSample, Person, utcnow

# 中文表头 → 内部字段；「服务状态」映射为报名状态
HEADER_MAP = {
    "姓名": "name",
    "手机": "phone",
    "所在城市": "city",
    "学院": "college",
    "学历": "education",
    "入学年份": "enroll_year",
    "身份识别照片": "photo_url",
    "票种": "ticket_type",
    "服务状态": "signup_status",
    "报名状态": "signup_status",
}


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    face_ok: int = 0
    face_failed: int = 0
    skipped: list[dict[str, Any]] = field(default_factory=list)


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def parse_xlsx_rows(data: bytes) -> list[dict[str, str]]:
    """解析 xlsx：在前 20 行找表头，返回数据行字典列表。"""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("空表格")

    header_idx = -1
    col_index: dict[str, int] = {}
    for i, row in enumerate(rows[:20]):
        cells = [_cell_str(c) for c in (row or ())]
        found: dict[str, int] = {}
        for j, text in enumerate(cells):
            key = HEADER_MAP.get(text)
            if key and key not in found:
                found[key] = j
        # 至少要有姓名和手机两列才认定表头
        if "name" in found and "phone" in found:
            header_idx = i
            col_index = found
            break

    if header_idx < 0:
        raise ValueError("未找到表头行（需包含「姓名」「手机」等列）")

    out: list[dict[str, str]] = []
    for row in rows[header_idx + 1 :]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        item = {field: "" for field in HEADER_MAP.values()}
        # 去重字段名（服务状态/报名状态都映射 signup_status）
        item = {
            "name": "",
            "phone": "",
            "city": "",
            "college": "",
            "education": "",
            "enroll_year": "",
            "photo_url": "",
            "ticket_type": "",
            "signup_status": "",
        }
        for field_name, j in col_index.items():
            if j < len(row):
                item[field_name] = _cell_str(row[j])
        out.append(item)
    return out


def download_image(url: str, timeout: float = 30.0) -> bytes:
    """下载身份识别照片。"""
    with httpx.Client(timeout=timeout, follow_redirects=True) as client:
        resp = client.get(url)
        resp.raise_for_status()
        content_type = resp.headers.get("content-type", "")
        if content_type and "image" not in content_type and "octet-stream" not in content_type:
            # 部分 OSS 不返回 image/*，仍尝试解码
            pass
        return resp.content


def enroll_face_bytes(person: Person, raw: bytes, db: Session) -> None:
    """对人员录入一张人脸图；成功更新索引与 face_status。失败抛 FaceEngineError。"""
    enabled_count = sum(1 for s in person.samples if s.enabled)
    if enabled_count >= MAX_SAMPLES_PER_PERSON:
        raise FaceEngineError("too_many_samples", f"每人最多 {MAX_SAMPLES_PER_PERSON} 张启用样本")

    image = face_engine.decode_image(raw)
    feature, enroll_image = face_engine.extract_for_enroll(image)
    hit, _ = gallery_index.search(
        feature.embedding,
        threshold=MATCH_THRESHOLD,
        exclude_person_id=person.id,
    )
    if hit is not None:
        raise FaceEngineError(
            "duplicate_face",
            f"疑似与已有人员重复: {hit.phone} {hit.name} (score={hit.score:.3f})",
        )

    rel_dir = Path("persons") / str(person.id)
    dest_dir = UPLOAD_DIR / rel_dir
    dest_dir.mkdir(parents=True, exist_ok=True)
    rel_path = (rel_dir / f"{uuid.uuid4().hex}.jpg").as_posix()
    (UPLOAD_DIR / rel_path).write_bytes(encode_image_jpeg(enroll_image))

    sample = FaceSample(
        person_id=person.id,
        image_path=rel_path,
        embedding=embedding_to_bytes(feature.embedding),
        quality_score=feature.quality_score,
        bbox_json=bbox_to_json(feature.bbox),
        enabled=True,
    )
    db.add(sample)
    db.flush()
    gallery_index.upsert_sample(
        sample_id=sample.id,
        person_id=person.id,
        phone=person.phone,
        name=person.name,
        embedding=feature.embedding,
        person_active=person.status == "active",
        sample_enabled=True,
    )
    person.face_status = "ok"
    person.face_message = ""


def import_people_from_xlsx(data: bytes, db: Session) -> ImportResult:
    """解析并导入；手机为空跳过；有照片则尝试录脸。"""
    rows = parse_xlsx_rows(data)
    result = ImportResult()
    print(f"[import] rows={len(rows)}")

    for row in rows:
        phone = row.get("phone", "").replace(" ", "")
        name = row.get("name", "").strip()
        if not phone:
            result.skipped.append({"reason": "空手机", "name": name})
            print(f"[import] skip 空手机 name={name!r}")
            continue
        if not name:
            result.skipped.append({"reason": "空姓名", "phone": phone})
            print(f"[import] skip 空姓名 phone={phone}")
            continue

        person = db.scalar(
            select(Person).options(selectinload(Person.samples)).where(Person.phone == phone)
        )
        photo_url = row.get("photo_url", "").strip()
        is_new = person is None

        if is_new:
            person = Person(
                phone=phone,
                name=name,
                city=row.get("city", ""),
                college=row.get("college", ""),
                education=row.get("education", ""),
                enroll_year=row.get("enroll_year", ""),
                ticket_type=row.get("ticket_type", ""),
                signup_status=row.get("signup_status", ""),
                photo_url=photo_url,
                face_status="none",
                face_message="",
                status="active",
            )
            db.add(person)
            db.flush()
            person.samples = []
            result.created += 1
            print(f"[import] create {phone} {name}")
        else:
            person.name = name
            person.city = row.get("city", "")
            person.college = row.get("college", "")
            person.education = row.get("education", "")
            person.enroll_year = row.get("enroll_year", "")
            person.ticket_type = row.get("ticket_type", "")
            person.signup_status = row.get("signup_status", "")
            person.updated_at = utcnow()
            result.updated += 1
            print(f"[import] update {phone} {name}")

        # 录脸：无 URL → missing；有 URL 且（新建 / URL 变化 / 尚无成功脸）则尝试
        need_enroll = False
        if not photo_url:
            if person.face_status != "ok":
                person.photo_url = ""
                person.face_status = "missing"
                person.face_message = "无身份识别照片"
                print(f"[import] face missing {phone}")
        else:
            url_changed = photo_url != (person.photo_url or "")
            person.photo_url = photo_url
            if is_new or url_changed or person.face_status != "ok":
                need_enroll = True

        if need_enroll:
            try:
                print(f"[import] download {phone} {photo_url[:80]}")
                raw = download_image(photo_url)
                # 重新加载 samples 关系
                db.refresh(person, attribute_names=["samples"])
                enroll_face_bytes(person, raw, db)
                result.face_ok += 1
                print(f"[import] face ok {phone}")
            except FaceEngineError as exc:
                person.face_status = "failed"
                person.face_message = exc.message[:500]
                result.face_failed += 1
                print(f"[import] face fail {phone}: {exc.message}")
            except Exception as exc:  # noqa: BLE001
                person.face_status = "failed"
                person.face_message = f"下载或处理失败: {exc}"[:500]
                result.face_failed += 1
                print(f"[import] face fail {phone}: {exc}")

        db.commit()

    print(
        f"[import] done created={result.created} updated={result.updated} "
        f"face_ok={result.face_ok} face_failed={result.face_failed} skipped={len(result.skipped)}"
    )
    return result
